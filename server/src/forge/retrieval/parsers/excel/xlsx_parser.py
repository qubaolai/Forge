"""XLSX parser for knowledge-base ingestion."""

from __future__ import annotations

import contextlib
import logging
from itertools import zip_longest
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from forge.core.types import Element
from forge.retrieval.parsers.parser_base import BaseParser

from .common import elements_from_tabular_rows

logger = logging.getLogger(__name__)


class XlsxParser(BaseParser):
    """Parse .xlsx workbooks into sheet metadata and row elements."""

    SUPPORTED_EXTENSIONS = [".xlsx"]

    def parse(self, file_path: Path) -> list[Element]:
        value_workbook = load_workbook(file_path, read_only=True, data_only=True)
        formula_workbook = load_workbook(file_path, read_only=True, data_only=False)
        try:
            elements: list[Element] = []
            for worksheet in value_workbook.worksheets:
                formula_sheet = formula_workbook[worksheet.title]
                # read_only 模式下 openpyxl 信任文件里缓存的 <dimension>, 部分导出工具
                # 会写成错误的 A1, 导致只读到左上角一个单元格. reset 强制全量扫描.
                _reset_dimensions(worksheet)
                _reset_dimensions(formula_sheet)
                rows = _merged_rows(worksheet, formula_sheet)
                elements.extend(
                    elements_from_tabular_rows(
                        source="xlsx",
                        sheet_name=worksheet.title,
                        rows=rows,
                    )
                )
            logger.info(
                "XLSX 解析 %s: %d 工作表, %d 元素",
                file_path.name,
                len(value_workbook.worksheets),
                len(elements),
            )
            return elements
        finally:
            value_workbook.close()
            formula_workbook.close()


def _reset_dimensions(worksheet) -> None:
    """强制 openpyxl 忽略缓存的 <dimension>, 全量扫描单元格 (兼容旧版本)."""
    reset = getattr(worksheet, "reset_dimensions", None)
    if callable(reset):
        with contextlib.suppress(Exception):
            reset()


def _merged_rows(value_sheet, formula_sheet) -> list[tuple[int, list[Any]]]:
    """Prefer cached values, fallback to formula text when cache is empty."""
    rows: list[tuple[int, list[Any]]] = []
    value_rows = value_sheet.iter_rows(values_only=True)
    formula_rows = formula_sheet.iter_rows(values_only=True)
    for row_index, (value_row, formula_row) in enumerate(
        zip_longest(value_rows, formula_rows, fillvalue=()),
        start=1,
    ):
        merged: list[Any] = []
        width = max(len(value_row), len(formula_row))
        for index in range(width):
            value = value_row[index] if index < len(value_row) else None
            formula_value = formula_row[index] if index < len(formula_row) else None
            merged.append(
                formula_value
                if _is_blank(value) and not _is_blank(formula_value)
                else value
            )
        rows.append((row_index, merged))
    return rows


def _is_blank(value: Any) -> bool:
    return value is None or value == ""
